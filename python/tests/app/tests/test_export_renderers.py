"""
Tests for the export renderers (``habarico.export.renderers``).

The renderers are plain DRF renderer classes and do not need the database - all tests here
work with in-memory row dicts, exactly as the export endpoint (HAB-532) will hand them over.
"""

import codecs
import csv
import datetime
import decimal
import io
import sys
import uuid

import pytest
from django.utils import timezone, translation
from django.utils.translation import gettext_lazy as _
from habarico.export.renderers import (
    EXPORT_COLUMNS_CONTEXT_KEY,
    EXPORT_TITLE_CONTEXT_KEY,
    BaseExportRenderer,
    CsvRenderer,
    ExportColumn,
    XlsxRenderer,
)

ROWS = [
    {"title": "First", "status": "draft"},
    {"title": "Second", "status": "published"},
]


def render_csv(data, renderer: CsvRenderer | None = None, **renderer_context) -> bytes:
    """Render ``data`` with a (optionally customized) CSV renderer."""
    renderer = renderer if renderer is not None else CsvRenderer()
    return renderer.render(data, renderer_context=renderer_context)


def csv_rows(data, renderer: CsvRenderer | None = None, **renderer_context) -> list[list[str]]:
    """Render ``data`` to CSV and parse it back into a list of rows."""
    output = render_csv(data, renderer=renderer, **renderer_context)
    text = output.decode("utf-8-sig")
    delimiter = renderer.delimiter if renderer is not None else CsvRenderer.delimiter
    return list(csv.reader(io.StringIO(text), delimiter=delimiter))


def xlsx_sheet(data, renderer: XlsxRenderer | None = None, **renderer_context):
    """Render ``data`` to XLSX and load the resulting worksheet back with openpyxl."""
    from openpyxl import load_workbook

    renderer = renderer if renderer is not None else XlsxRenderer()
    output = renderer.render(data, renderer_context=renderer_context)
    return load_workbook(io.BytesIO(output)).active


class TestExportColumn:
    """Column descriptor and its normalization from the shorthand notations."""

    def test_header_falls_back_to_name(self):
        assert ExportColumn("title").header == "title"

    def test_header_uses_label(self):
        assert ExportColumn("title", label="Title").header == "Title"

    def test_header_resolves_lazy_label(self):
        assert ExportColumn("title", label=_("Yes")).header == "Yes"

    def test_normalize_accepts_names(self):
        assert BaseExportRenderer.normalize_columns(["title", "status"]) == [
            ExportColumn("title"),
            ExportColumn("status"),
        ]

    def test_normalize_accepts_pairs(self):
        assert BaseExportRenderer.normalize_columns([("title", "Title")]) == [ExportColumn("title", label="Title")]

    def test_normalize_accepts_mapping(self):
        assert BaseExportRenderer.normalize_columns({"title": "Title"}) == [ExportColumn("title", label="Title")]

    def test_normalize_keeps_columns(self):
        columns = [ExportColumn("title", label="Title")]
        assert BaseExportRenderer.normalize_columns(columns) == columns

    def test_normalize_returns_none_for_none(self):
        assert BaseExportRenderer.normalize_columns(None) is None


class TestCsvRendererStructure:
    """Header, column selection and the accepted shapes of the rendered data."""

    def test_header_comes_from_columns(self):
        rows = csv_rows(ROWS, **{EXPORT_COLUMNS_CONTEXT_KEY: [ExportColumn("title", label="Title")]})
        assert rows == [["Title"], ["First"], ["Second"]]

    def test_header_falls_back_to_row_keys(self):
        assert csv_rows(ROWS) == [["title", "status"], ["First", "draft"], ["Second", "published"]]

    def test_columns_accept_plain_names(self):
        assert csv_rows(ROWS, **{EXPORT_COLUMNS_CONTEXT_KEY: ["status"]}) == [["status"], ["draft"], ["published"]]

    def test_paginated_data_is_unwrapped(self):
        data = {"count": 2, "results": ROWS}
        assert csv_rows(data) == [["title", "status"], ["First", "draft"], ["Second", "published"]]

    def test_single_object_is_rendered_as_one_row(self):
        assert csv_rows(ROWS[0]) == [["title", "status"], ["First", "draft"]]

    def test_rows_may_be_a_generator(self):
        assert csv_rows(row for row in ROWS) == [["title", "status"], ["First", "draft"], ["Second", "published"]]

    def test_missing_key_renders_empty_cell(self):
        rows = csv_rows([{"title": "First"}], **{EXPORT_COLUMNS_CONTEXT_KEY: ["title", "status"]})
        assert rows == [["title", "status"], ["First", ""]]

    def test_nested_value_is_read_through_lookup_path(self):
        data = [{"author": {"name": "Jan"}}]
        assert csv_rows(data, **{EXPORT_COLUMNS_CONTEXT_KEY: ["author__name"]}) == [["author__name"], ["Jan"]]

    def test_empty_rows_render_header_only(self):
        rows = csv_rows([], **{EXPORT_COLUMNS_CONTEXT_KEY: ["title"]})
        assert rows == [["title"]]

    def test_no_rows_and_no_columns_render_nothing(self):
        assert render_csv([]) == codecs.BOM_UTF8


class TestCsvRendererValues:
    """Value formatting shared by all export renderers, in its CSV (text) form."""

    def _cell(self, value, renderer: CsvRenderer | None = None) -> str:
        return csv_rows([{"value": value}], renderer=renderer)[1][0]

    def test_none_is_empty(self):
        assert self._cell(None) == ""

    def test_booleans_use_localized_labels(self):
        assert self._cell(True) == "Yes"
        assert self._cell(False) == "No"

    def test_booleans_are_translated(self):
        with translation.override("cs"):
            assert self._cell(True) == "Ano"

    def test_lazy_string_is_resolved(self):
        assert self._cell(_("Yes")) == "Yes"

    def test_numbers_keep_their_text_form(self):
        assert self._cell(42) == "42"
        assert self._cell(decimal.Decimal("10.50")) == "10.50"

    def test_dates_are_iso_formatted(self):
        assert self._cell(datetime.date(2026, 7, 29)) == "2026-07-29"
        assert self._cell(datetime.datetime(2026, 7, 29, 8, 30)) == "2026-07-29T08:30:00"
        assert self._cell(datetime.time(8, 30)) == "08:30:00"

    def test_uuid_is_stringified(self):
        value = uuid.uuid4()
        assert self._cell(value) == str(value)

    def test_list_is_joined(self):
        assert self._cell(["a", "b"]) == "a, b"

    def test_dict_uses_its_label(self):
        assert self._cell({"id": 1, "label": "Jan"}) == "Jan"

    def test_dict_without_label_falls_back_to_json(self):
        assert self._cell({"id": 1, "code": "x"}) == '{"id": 1, "code": "x"}'

    def test_delimiter_and_newline_are_quoted(self):
        rows = csv_rows([{"value": "a,b\nc"}])
        assert rows[1] == ["a,b\nc"]


class TestCsvRendererFormat:
    """Dialect, encoding and the file naming helper."""

    def test_uses_crlf_line_terminator(self):
        assert render_csv(ROWS).endswith(b"published\r\n")

    def test_bom_is_written_by_default(self):
        assert render_csv(ROWS).startswith(codecs.BOM_UTF8)

    def test_bom_can_be_disabled(self):
        class NoBomCsvRenderer(CsvRenderer):
            use_bom = False

        assert not render_csv(ROWS, renderer=NoBomCsvRenderer()).startswith(codecs.BOM_UTF8)

    def test_delimiter_can_be_overridden(self):
        class SemicolonCsvRenderer(CsvRenderer):
            delimiter = ";"

        output = render_csv(ROWS, renderer=SemicolonCsvRenderer())
        assert b"title;status" in output

    def test_non_ascii_is_utf8_encoded(self):
        assert "Crème brûlée" in render_csv([{"value": "Crème brûlée"}]).decode("utf-8-sig")

    def test_write_accepts_a_stream(self):
        stream = io.BytesIO()
        CsvRenderer().write(ROWS, stream, columns=["title"])
        assert stream.getvalue().decode("utf-8-sig").splitlines() == ["title", "First", "Second"]

    def test_write_only_needs_a_write_method(self):
        """The asynchronous runner writes into a Django file object, not into a full stream."""

        class Sink:
            def __init__(self):
                self.chunks = []

            def write(self, data):
                self.chunks.append(data)

        sink = Sink()
        CsvRenderer().write(ROWS, sink, columns=["title"])
        assert b"".join(sink.chunks).decode("utf-8-sig").splitlines() == ["title", "First", "Second"]
        assert len(sink.chunks) == 3  # header and each row flushed separately

    def test_get_filename_appends_the_extension(self):
        assert CsvRenderer().get_filename("article") == "article.csv"

    def test_media_type_and_format(self):
        assert CsvRenderer.media_type == "text/csv"
        assert CsvRenderer.format == "csv"


class TestFormulaEscaping:
    """Values that spreadsheet applications would otherwise evaluate as formulas."""

    @pytest.mark.parametrize("value", ["=1+1", "+1", "-1+1", "@SUM(A1)", "\tx", "\rx"])
    def test_dangerous_values_are_prefixed(self, value):
        assert csv_rows([{"value": value}])[1][0] == f"'{value}"

    def test_plain_values_are_untouched(self):
        assert csv_rows([{"value": "plain"}])[1][0] == "plain"

    def test_negative_numbers_are_not_escaped(self):
        assert csv_rows([{"value": -1}])[1][0] == "-1"
        assert csv_rows([{"value": decimal.Decimal("-1.5")}])[1][0] == "-1.5"

    def test_escaping_can_be_disabled(self):
        class RawCsvRenderer(CsvRenderer):
            escape_formulas = False

        assert csv_rows([{"value": "=1+1"}], renderer=RawCsvRenderer())[1][0] == "=1+1"


class TestXlsxRenderer:
    """XLSX output read back with openpyxl."""

    def test_header_and_values(self):
        sheet = xlsx_sheet(ROWS, **{EXPORT_COLUMNS_CONTEXT_KEY: [ExportColumn("title", label="Title")]})
        assert [[cell.value for cell in row] for row in sheet.rows] == [["Title"], ["First"], ["Second"]]

    def test_header_falls_back_to_row_keys(self):
        sheet = xlsx_sheet(ROWS)
        assert [cell.value for cell in next(sheet.rows)] == ["title", "status"]

    def test_numbers_stay_numeric(self):
        sheet = xlsx_sheet([{"value": 42, "price": decimal.Decimal("10.50")}])
        row = list(sheet.rows)[1]
        assert row[0].value == 42
        assert row[1].value == decimal.Decimal("10.50")

    def test_dates_stay_native(self):
        sheet = xlsx_sheet([{"value": datetime.date(2026, 7, 29)}])
        assert list(sheet.rows)[1][0].value == datetime.datetime(2026, 7, 29)

    def test_aware_datetime_is_converted_to_local_naive(self):
        value = datetime.datetime(2026, 7, 29, 8, 30, tzinfo=datetime.timezone.utc)
        sheet = xlsx_sheet([{"value": value}])
        assert list(sheet.rows)[1][0].value == timezone.localtime(value).replace(tzinfo=None)

    def test_booleans_use_localized_labels(self):
        sheet = xlsx_sheet([{"value": True}])
        assert list(sheet.rows)[1][0].value == "Yes"

    def test_none_is_empty(self):
        sheet = xlsx_sheet([{"value": None, "other": "x"}])
        assert [cell.value for cell in list(sheet.rows)[1]] == [None, "x"]

    def test_header_is_bold_and_frozen(self):
        sheet = xlsx_sheet(ROWS)
        assert next(sheet.rows)[0].font.bold is True
        assert sheet.freeze_panes == "A2"

    def test_illegal_characters_are_stripped(self):
        sheet = xlsx_sheet([{"value": "bad\x0bvalue"}])
        assert list(sheet.rows)[1][0].value == "badvalue"

    def test_long_value_is_truncated(self):
        sheet = xlsx_sheet([{"value": "x" * 40000}])
        assert len(list(sheet.rows)[1][0].value) == XlsxRenderer.max_cell_length

    def test_formula_is_escaped(self):
        sheet = xlsx_sheet([{"value": "=1+1"}])
        assert list(sheet.rows)[1][0].value == "'=1+1"

    def test_sheet_title_comes_from_the_context(self):
        assert xlsx_sheet(ROWS, **{EXPORT_TITLE_CONTEXT_KEY: "Articles"}).title == "Articles"

    def test_sheet_title_is_sanitized_and_shortened(self):
        sheet = xlsx_sheet(ROWS, **{EXPORT_TITLE_CONTEXT_KEY: "Invalid[]:*?/\\" + "x" * 40})
        assert len(sheet.title) <= 31
        assert not set(sheet.title) & set("[]:*?/\\")

    def test_empty_title_falls_back_to_the_default(self):
        assert xlsx_sheet(ROWS, **{EXPORT_TITLE_CONTEXT_KEY: "  "}).title == str(XlsxRenderer.sheet_title)

    def test_write_accepts_a_stream(self):
        from openpyxl import load_workbook

        stream = io.BytesIO()
        XlsxRenderer().write(ROWS, stream, columns=["title"])
        assert [cell.value for cell in load_workbook(stream).active["A"]] == ["title", "First", "Second"]

    def test_get_filename_appends_the_extension(self):
        assert XlsxRenderer().get_filename("article") == "article.xlsx"

    def test_media_type_and_format(self):
        assert XlsxRenderer.format == "xlsx"
        assert XlsxRenderer.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert XlsxRenderer.charset is None

    def test_missing_openpyxl_raises_a_helpful_error(self, monkeypatch):
        monkeypatch.setitem(sys.modules, "openpyxl", None)
        with pytest.raises(ImportError, match=r"export-xlsx"):
            XlsxRenderer().render(ROWS)
